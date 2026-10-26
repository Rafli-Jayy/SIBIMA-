"""Aturan akses, penyusunan laporan admin, serta ekspor CSV/XLSX."""

import csv
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path

import psycopg

from app import session
from app.config import BASE_DIR
from app.repositories import log_repo, report_repo, user_repo
from app.utils.csv_safety import nilai_csv_aman

LAPORAN = {
    "pemasukan": {
        "judul": "Pemasukan per Bulan",
        "kolom": [
            ("bulan", "Bulan"),
            ("jumlah_transaksi", "Jumlah transaksi lunas"),
            ("total_pemasukan", "Total pemasukan (Rp)"),
        ],
        "ambil": report_repo.pemasukan_bulanan,
    },
    "siswa_program": {
        "judul": "Siswa per Program",
        "kolom": [
            ("kode_program", "Kode program"), ("nama_program", "Nama program"),
            ("jenjang", "Jenjang"), ("status_program", "Status program"),
            ("kuota", "Kuota"), ("siswa_aktif", "Siswa aktif"),
            ("menunggu", "Menunggu persetujuan"), ("kuota_tersisa", "Sisa kuota"),
        ],
        "ambil": report_repo.siswa_per_program,
    },
    "kehadiran": {
        "judul": "Rekap Kehadiran",
        "kolom": [
            ("kode_program", "Kode program"), ("nama_program", "Nama program"),
            ("nis", "NIS"), ("nama_siswa", "Nama siswa"),
            ("total_pertemuan", "Total pertemuan"), ("hadir", "Hadir"),
            ("izin", "Izin"), ("sakit", "Sakit"), ("alpa", "Alpa"),
            ("persen_hadir", "Kehadiran (%)"),
        ],
        "ambil": report_repo.rekap_kehadiran,
    },
    "nilai": {
        "judul": "Rata-rata Nilai per Program",
        "kolom": [
            ("kode_program", "Kode program"), ("nama_program", "Nama program"),
            ("jenjang", "Jenjang"), ("siswa_dinilai", "Siswa sudah dinilai"),
            ("rata_nilai_akhir", "Rata-rata nilai akhir"),
        ],
        "ambil": report_repo.rata_nilai_per_program,
    },
}


def _admin_aktif() -> dict:
    pengguna = session.get_user() or {}
    if pengguna.get("role") != "admin" or type(pengguna.get("id")) is not int:
        raise PermissionError("Fitur laporan hanya untuk admin aktif.")
    admin = user_repo.cari_by_id(pengguna["id"])
    if not admin or admin.get("role") != "admin" or admin.get("status") != "aktif":
        raise PermissionError("Akun admin tidak aktif atau sesi tidak valid.")
    return admin


def daftar_jenis_laporan() -> list[dict]:
    _admin_aktif()
    return [{"kode": kode, "judul": item["judul"]} for kode, item in LAPORAN.items()]


def ambil_laporan(kode: str) -> dict:
    _admin_aktif()
    if kode not in LAPORAN:
        raise ValueError("Jenis laporan tidak dikenal.")
    definisi = LAPORAN[kode]
    data = definisi["ambil"]()
    return {"kode": kode, "judul": definisi["judul"], "kolom": definisi["kolom"], "data": data}


def _siapkan_data_ekspor(kode: str):
    admin = _admin_aktif()
    if kode not in LAPORAN:
        raise ValueError("Jenis laporan tidak dikenal.")
    definisi = LAPORAN[kode]
    data = definisi["ambil"]()
    return admin, definisi, data


def _direktori_laporan() -> Path:
    direktori = Path(BASE_DIR) / "laporan"
    direktori.mkdir(parents=True, exist_ok=True)
    return direktori


def _catat_ekspor(admin: dict, kode_aksi: str, definisi: dict, path: Path, jumlah: int) -> None:
    try:
        log_repo.catat(
            admin["id"],
            kode_aksi,
            f"Laporan {definisi['judul']} diekspor ke {path.name} ({jumlah} baris).",
        )
    except psycopg.Error:
        # File hasil ekspor tetap valid walau pencatatan log sedang bermasalah.
        pass


def ekspor_csv(kode: str) -> Path:
    """Ekspor CSV UTF-8-BOM dengan pemisah titik koma yang ramah Excel Indonesia."""
    admin, definisi, data = _siapkan_data_ekspor(kode)
    kolom = definisi["kolom"]
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S_%f")
    path = _direktori_laporan() / f"{kode}_{timestamp}.csv"

    try:
        with path.open("w", newline="", encoding="utf-8-sig") as file:
            writer = csv.writer(file, delimiter=";", lineterminator="\r\n", quoting=csv.QUOTE_MINIMAL)
            writer.writerow([label for _, label in kolom])
            for baris in data:
                writer.writerow([nilai_csv_aman(baris.get(key)) for key, _ in kolom])
    except PermissionError as exc:
        raise OSError("File CSV tidak bisa ditulis. Tutup file yang mungkin masih terbuka, lalu coba lagi.") from exc

    _catat_ekspor(admin, "LAPORAN_DIEKSPOR", definisi, path, len(data))
    return path


def ekspor_excel(kode: str) -> Path:
    """Ekspor XLSX dengan judul, format angka, filter, freeze pane, dan total pemasukan."""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
        from openpyxl.utils import get_column_letter
        from openpyxl.worksheet.page import PageMargins
    except ImportError as exc:
        raise RuntimeError("Paket openpyxl belum terpasang. Jalankan: python -m pip install 'openpyxl>=3.1'.") from exc

    admin, definisi, data = _siapkan_data_ekspor(kode)
    kolom = definisi["kolom"]
    jumlah_kolom = len(kolom)
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S_%f")
    path = _direktori_laporan() / f"{kode}_{timestamp}.xlsx"

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = definisi["judul"][:31]
    workbook.properties.creator = "SIBIMA"
    workbook.properties.title = f"Laporan SIBIMA - {definisi['judul']}"
    workbook.properties.subject = "Laporan administrasi akademik"
    sheet.sheet_view.showGridLines = False
    sheet.sheet_properties.tabColor = "173B68"

    navy = "173B68"
    blue = "24578A"
    pale_blue = "EAF1F8"
    stripe = "F5F8FC"
    green = "E2F0D9"
    border_color = "D7E0EA"
    white = "FFFFFF"
    thin_bottom = Border(bottom=Side(style="thin", color=border_color))

    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=jumlah_kolom)
    title_cell = sheet.cell(1, 1, f"SIBIMA  |  {definisi['judul'].upper()}")
    title_cell.fill = PatternFill("solid", fgColor=navy)
    title_cell.font = Font(name="Aptos Display", size=16, bold=True, color=white)
    title_cell.alignment = Alignment(vertical="center", horizontal="left")
    sheet.row_dimensions[1].height = 32
    for col in range(1, jumlah_kolom + 1):
        sheet.cell(1, col).fill = PatternFill("solid", fgColor=navy)

    sheet.merge_cells(start_row=2, start_column=1, end_row=2, end_column=jumlah_kolom)
    subtitle = sheet.cell(2, 1, f"Dihasilkan: {datetime.now().strftime('%d %B %Y, %H:%M')}  |  Sumber: database SIBIMA")
    subtitle.fill = PatternFill("solid", fgColor=pale_blue)
    subtitle.font = Font(name="Aptos", size=10, color="40536A", italic=True)
    subtitle.alignment = Alignment(vertical="center")
    sheet.row_dimensions[2].height = 22
    for col in range(1, jumlah_kolom + 1):
        sheet.cell(2, col).fill = PatternFill("solid", fgColor=pale_blue)

    header_row = 4
    for nomor, (_, label) in enumerate(kolom, start=1):
        cell = sheet.cell(header_row, nomor, label)
        cell.fill = PatternFill("solid", fgColor=blue)
        cell.font = Font(name="Aptos", size=10, bold=True, color=white)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin_bottom
    sheet.row_dimensions[header_row].height = 30

    start_data = header_row + 1
    for row_index, baris in enumerate(data, start=start_data):
        for col_index, (key, _) in enumerate(kolom, start=1):
            value = baris.get(key)
            if isinstance(value, Decimal):
                value = float(value)
            cell = sheet.cell(row_index, col_index, value)
            cell.font = Font(name="Aptos", size=10, color="203047")
            cell.border = thin_bottom
            cell.alignment = Alignment(vertical="center", horizontal="right" if isinstance(value, (int, float)) else "left")
            if (row_index - start_data) % 2:
                cell.fill = PatternFill("solid", fgColor=stripe)
            if key == "bulan" and isinstance(value, (date, datetime)):
                cell.number_format = "mmmm yyyy"
            elif key == "total_pemasukan":
                cell.number_format = '"Rp" #,##0;[Red]-"Rp" #,##0'
            elif key == "persen_hadir":
                cell.number_format = '0.0"%"'
            elif key == "rata_nilai_akhir":
                cell.number_format = "0.00"
            elif key in {"jumlah_transaksi", "kuota", "siswa_aktif", "menunggu", "kuota_tersisa", "total_pertemuan", "hadir", "izin", "sakit", "alpa", "siswa_dinilai"}:
                cell.number_format = "#,##0"
        sheet.row_dimensions[row_index].height = 22

    last_data_row = start_data + len(data) - 1
    if not data:
        sheet.merge_cells(start_row=start_data, start_column=1, end_row=start_data, end_column=jumlah_kolom)
        empty_cell = sheet.cell(start_data, 1, "Tidak ada data untuk laporan ini.")
        empty_cell.font = Font(name="Aptos", italic=True, color="66758A")
        empty_cell.alignment = Alignment(horizontal="center")
        last_data_row = header_row

    if data:
        sheet.auto_filter.ref = f"A{header_row}:{get_column_letter(jumlah_kolom)}{last_data_row}"
    else:
        sheet.auto_filter.ref = f"A{header_row}:{get_column_letter(jumlah_kolom)}{header_row}"

    # Total hanya dijumlahkan pada laporan pemasukan; data siswa/kehadiran tidak dijumlahkan
    # untuk menghindari perhitungan ganda pada siswa yang terdaftar di lebih dari satu program.
    if kode == "pemasukan" and data:
        total_row = last_data_row + 1
        sheet.cell(total_row, 1, "TOTAL PADA DATA YANG DITAMPILKAN")
        sheet.cell(total_row, 2, sum(int(row.get("jumlah_transaksi") or 0) for row in data))
        total_nominal = sum((Decimal(str(row.get("total_pemasukan") or 0)) for row in data), Decimal("0"))
        total_value = sheet.cell(total_row, 3, float(total_nominal))
        total_value.number_format = '"Rp" #,##0;[Red]-"Rp" #,##0'
        for col in range(1, jumlah_kolom + 1):
            cell = sheet.cell(total_row, col)
            cell.fill = PatternFill("solid", fgColor=green)
            cell.font = Font(name="Aptos", bold=True, color="203047")
            cell.border = Border(top=Side(style="medium", color="6A9955"))
            cell.alignment = Alignment(vertical="center", horizontal="right" if col > 1 else "left", wrap_text=True)
        sheet.row_dimensions[total_row].height = 26

    # Lebar kolom adaptif; batasi agar lembar tetap mudah dibaca dan dicetak.
    for col_index, (key, label) in enumerate(kolom, start=1):
        sample_values = [str(row.get(key) if row.get(key) is not None else "") for row in data[:200]]
        max_len = max([len(label), *(len(v) for v in sample_values)], default=len(label))
        width = min(max(max_len + 3, 12), 42)
        if key in {"nama_program", "nama_siswa"}:
            width = min(max(width, 24), 36)
        elif key == "total_pemasukan":
            width = 22
        sheet.column_dimensions[get_column_letter(col_index)].width = width

    sheet.freeze_panes = "A5"
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = "landscape" if jumlah_kolom >= 7 else "portrait"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.page_margins = PageMargins(left=0.25, right=0.25, top=0.5, bottom=0.5, header=0.2, footer=0.2)
    sheet.print_title_rows = "1:4"
    sheet.oddFooter.left.text = "SIBIMA - Laporan internal"
    sheet.oddFooter.right.text = "Halaman &P dari &N"

    try:
        workbook.save(path)
    except PermissionError as exc:
        raise OSError("File Excel tidak bisa ditulis. Tutup workbook yang mungkin masih terbuka, lalu coba lagi.") from exc

    _catat_ekspor(admin, "LAPORAN_DIEKSPOR_XLSX", definisi, path, len(data))
    return path
